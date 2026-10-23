"""
Daily driver.

    python run_daily.py              # pull stats + lines, project, write outputs/YYYY-MM-DD_props.xlsx
    python run_daily.py --close      # ONLY re-pull lines and save a closing snapshot (run ~15 min before first tip)
    python run_daily.py --demo       # synthetic data, no network — pipeline test
    python run_daily.py --no-ledger  # don't append tonight's picks to the ledger
    python run_daily.py --reuse-lines # re-run the model on today's already-saved lines (no API credits)
    python run_daily.py --rescore    # afternoon: re-pull injuries, re-score against the morning lines,
                                     #   write outputs/YYYY-MM-DD_props_HHMM.xlsx and print what changed

Typical day:
   morning:   python run_daily.py          -> bet from the "Bets" sheet
   pre-tip:   python run_daily.py --close  -> captures closing lines for CLV
   next day:  python grade.py              -> hit rate / ROI / calibration / CLV
"""
import os
import sys
if "--demo" in sys.argv:
    os.environ["LINES_SUBDIR"] = "lines_demo"
from datetime import date

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from config import LEDGER, OUT_DIR, SEASON, PRIOR_SEASON
from fetch_lines import devig
from model import blend_team_stats, build_projections, evaluate, pick_bets

LEDGER_COLS = ["date", "game", "player_id", "player", "player_key", "team", "opp", "stat", "side",
               "line", "book", "price", "projection", "model_prob", "mkt_fair_prob", "blended_prob",
               "edge", "ev_per_dollar", "stake", "n_cur", "dtd", "inj_note", "tier", "abs_spread", "event_id", "fetched_utc",
               "actual", "result", "close_price", "close_fair_prob", "clv", "profit"]


def main():
    args = set(sys.argv[1:])
    demo = "--demo" in args
    rescore = "--rescore" in args
    if rescore:
        args.add("--reuse-lines")
    today = date.today()

    if "--close" in args and not demo:
        from fetch_lines import fetch_lines
        fetch_lines("close")
        print("Closing snapshot saved. Run grade.py tomorrow.")
        return

    # ---- data ------------------------------------------------------------
    if demo:
        from demo_data import make_demo
        cur, prior, t_cur, t_prior, lines = make_demo(today)
        print("DEMO MODE: synthetic data.")
    else:
        from fetch_stats import load_all
        from fetch_lines import fetch_lines
        print(f"Pulling game logs {SEASON} + {PRIOR_SEASON} and team stats ...")
        cur, prior, t_cur, t_prior = load_all()
        from fetch_lines import load_snapshot
        lines = load_snapshot(today.isoformat(), "open") if "--reuse-lines" in args else None
        if lines is None:
            print("Pulling lines from The Odds API ...")
            lines = fetch_lines("open")
        else:
            print("Reusing today's saved open lines.")
    if lines.empty:
        print("No lines returned — no games today, or all games already tipped.")
        return

    # ---- model -------------------------------------------------------------
    lines = devig(lines)
    slate = lines[["home", "away"]].drop_duplicates()
    teams = blend_team_stats(t_cur, t_prior)
    if demo:
        rosters, inj = None, None
    else:
        from fetch_stats import rosters as _rosters
        from injuries import fetch_injuries
        rosters, inj = _rosters(), fetch_injuries()
        on_slate = inj[inj["team"].isin(set(slate["home"]) | set(slate["away"]))]
        print(f"Injury report: {int(on_slate['is_out'].sum())} OUT, {int(on_slate['is_dtd'].sum())} day-to-day on tonight's slate "
              f"({len(inj)} league-wide).")
    proj = build_projections(cur, prior, teams, slate, today, rosters, inj)
    print(f"{len(proj)} players projected across {len(slate)} games.")
    evald = evaluate(proj, lines)
    bets = pick_bets(evald)
    print(f"{len(evald)} priced lines evaluated, {len(bets)} qualify (edge floor {config.MIN_EDGE_UNDER:.1%} unders / {config.MIN_EDGE_OVER:.0%} overs).")

    # ---- outputs -----------------------------------------------------------
    from datetime import datetime
    suffix = "_demo" if demo else (f"_{datetime.now().strftime('%H%M')}" if rescore else "")
    out = OUT_DIR / f"{today.isoformat()}_props{suffix}.xlsx"
    if rescore:
        report_changes(bets, today)
    write_xlsx(out, bets, evald, proj, lines, inj)
    print(f"Wrote {out}")

    if "--no-ledger" not in args:
        if LEDGER.exists():   # re-run same day: replace today's picks (even with an empty board)
            old = pd.read_csv(LEDGER)
            old[old["date"] != today.isoformat()].to_csv(LEDGER, index=False)
        if not bets.empty:
            append_ledger(bets, today)
        print(f"Appended {len(bets)} picks to {LEDGER}")

    cols = ["player", "stat", "side", "line", "book", "price", "projection", "blended_prob", "edge", "stake", "n_cur"]
    if not bets.empty:
        print("\nTOP BETS (tier: blowout = under in a 10+ pt spread game, the strongest angle in the backtest; "
              "core = under, 10+ games, no injury-driven projection; over (experimental) = close game, model 5%+ "
              "above the line, opponent allows above average — half stakes, weak evidence; (early) = player has 2-4 games, half stakes)\n"
              + bets[["tier", "abs_spread"] + cols].head(20).to_string(index=False))
    else:
        watch = evald[(evald["edge"] >= config.MIN_EDGE_UNDER) & (~evald["suspicious"])
                      & (evald["book"].isin(config.BETTABLE_BOOKS) if config.BETTABLE_BOOKS else True)
                      ].sort_values("edge", ascending=False)
        watch = watch.drop_duplicates(["player_key", "stat"]).head(10)
        print(f"\nNO BETS qualify (every tier needs {config.EARLY_MIN_N_CUR}+ games this season; core tier = under with no usage-model involvement, "
              f"blowout tier = under in a {config.BLOWOUT_SPREAD:.0f}+ pt spread game; "
              f"experimental over tier = close game, projection 5%+ over the line, soft opponent). "
              f"Watchlist of what the model likes anyway — NOT recommendations:\n" + watch[cols].to_string(index=False))


def report_changes(bets: pd.DataFrame, today: date):
    """Compare tonight's new board with the most recent board written today."""
    prev_files = sorted(p for p in OUT_DIR.glob(f"{today.isoformat()}_props*.xlsx") if "demo" not in p.name)
    if not prev_files:
        return
    prev = pd.read_excel(prev_files[-1], "Bets")
    key = ["player", "stat", "side"]
    old = prev[key].apply(tuple, axis=1).tolist() if not prev.empty else []
    new = bets[key].apply(tuple, axis=1).tolist() if not bets.empty else []
    dropped = [k for k in old if k not in new]
    added = [k for k in new if k not in old]
    print(f"\nRESCORE vs {prev_files[-1].name}: {len(added)} bets added, {len(dropped)} dropped, "
          f"{len(set(old) & set(new))} unchanged.")
    if dropped:
        print("DROPPED (player scratched, teammate news, or edge gone):")
        for k in dropped[:25]:
            print("   -", *k)
    if added:
        print("ADDED:")
        for k in added[:25]:
            r = bets[(bets["player"] == k[0]) & (bets["stat"] == k[1]) & (bets["side"] == k[2])].iloc[0]
            print(f"   + {k[0]} {k[1]} {k[2]} {r['line']} @ {r['book']} {r['price']:+d} (edge {r['edge']:.1%}) {r.get('inj_note', '')}")


def append_ledger(bets: pd.DataFrame, today: date):
    rows = bets.copy()
    rows.insert(0, "date", today.isoformat())
    for c in LEDGER_COLS:
        if c not in rows.columns:
            rows[c] = pd.NA
    rows = rows[LEDGER_COLS]
    if LEDGER.exists():
        old = pd.read_csv(LEDGER)
        # don't double-log the same pick if you run twice in a day
        key = ["date", "player_key", "stat", "side"]
        rows = rows.merge(old[key].drop_duplicates(), on=key, how="left", indicator=True)
        rows = rows[rows["_merge"] == "left_only"].drop(columns="_merge")
        rows = pd.concat([old, rows], ignore_index=True)
    rows.to_csv(LEDGER, index=False)


def write_xlsx(path, bets, evald, proj, lines, inj=None):
    bet_cols = ["tier", "abs_spread", "game", "player", "team", "opp", "stat", "side", "line", "book", "price",
                "projection", "minutes", "model_prob", "mkt_fair_prob", "blended_prob", "edge",
                "ev_per_dollar", "fair_price", "stake", "n_books", "n_cur", "new_team", "b2b",
                "market_weight", "disagree", "suspicious", "dtd", "inj_status", "min_boost", "inj_note", "opp_allow_rel"]
    proj_cols = ["game", "player", "team", "opp", "pos", "home", "b2b", "minutes", "min_boost", "n_cur", "new_team",
                 "dtd", "inj_status", "inj_note", "mu_PTS", "mu_REB", "mu_AST", "mu_FG3M", "mu_PRA",
                 "posdef_PTS", "posdef_REB", "posdef_AST", "posdef_FG3M"]
    knobs = pd.DataFrame({"setting": ["MARKET_WEIGHT", "MIN_EDGE", "MIN_BOOKS", "KELLY_FRACTION",
                                      "MAX_BET_PCT", "BANKROLL", "PRIOR_SHRINK_K", "OPP_ADJ_STRENGTH"],
                          "value": [config.MARKET_WEIGHT, config.MIN_EDGE, config.MIN_BOOKS,
                                    config.KELLY_FRACTION, config.MAX_BET_PCT, config.BANKROLL,
                                    config.PRIOR_SHRINK_K, config.OPP_ADJ_STRENGTH]})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        (bets[[c for c in bet_cols if c in bets.columns]] if not bets.empty else pd.DataFrame(columns=bet_cols)
         ).to_excel(xw, sheet_name="Bets", index=False)
        evald[[c for c in bet_cols if c in evald.columns]].to_excel(xw, sheet_name="All lines", index=False)
        proj[[c for c in proj_cols if c in proj.columns]].sort_values(["game", "mu_PTS"], ascending=[True, False]
                                                                       ).to_excel(xw, sheet_name="Projections", index=False)
        lines.to_excel(xw, sheet_name="Raw lines", index=False)
        if inj is not None and not inj.empty:
            inj.drop(columns=["player_key"]).to_excel(xw, sheet_name="Injuries", index=False)
        knobs.to_excel(xw, sheet_name="Settings", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3A5F")
            for i, col in enumerate(ws.columns, 1):
                w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 28)
            ws.auto_filter.ref = ws.dimensions
            for name, fmt in (("model_prob", "0.0%"), ("mkt_fair_prob", "0.0%"), ("blended_prob", "0.0%"),
                              ("edge", "+0.0%"), ("ev_per_dollar", "+0.0%"), ("stake", "$#,##0.00")):
                for j, c in enumerate(ws[1], 1):
                    if c.value == name:
                        for cell in list(ws.columns)[j - 1][1:]:
                            cell.number_format = fmt


if __name__ == "__main__":
    main()
